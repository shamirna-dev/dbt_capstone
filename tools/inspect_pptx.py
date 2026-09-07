"""Dump the structure of the L2 capstone template so we know exactly what to fill."""
import sys
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1]
prs = Presentation(path)

print(f"SLIDE_SIZE: {prs.slide_width} x {prs.slide_height} "
      f"({Emu(prs.slide_width).inches:.2f} x {Emu(prs.slide_height).inches:.2f} in)")
print(f"SLIDES: {len(prs.slides)}")
print("=" * 78)

for i, slide in enumerate(prs.slides, start=1):
    layout = slide.slide_layout.name
    print(f"\n--- SLIDE {i}  (layout: {layout}) ---")
    for shape in slide.shapes:
        kind = shape.shape_type
        name = shape.name
        ph = ""
        if shape.is_placeholder:
            ph = f" PLACEHOLDER idx={shape.placeholder_format.idx} type={shape.placeholder_format.type}"
        pos = (f" @({Emu(shape.left).inches:.2f},{Emu(shape.top).inches:.2f}) "
               f"{Emu(shape.width).inches:.2f}x{Emu(shape.height).inches:.2f}in"
               if shape.left is not None else "")
        print(f"  [{name}] {kind}{ph}{pos}")
        if shape.has_text_frame:
            for p_i, para in enumerate(shape.text_frame.paragraphs):
                txt = "".join(r.text for r in para.runs)
                if txt.strip():
                    print(f"      p{p_i} (lvl{para.level}): {txt}")
        if shape.has_table:
            tbl = shape.table
            print(f"      TABLE {len(tbl.rows)}x{len(tbl.columns)}")
            for r_i, row in enumerate(tbl.rows):
                cells = [c.text.replace("\n", " / ") for c in row.cells]
                print(f"        r{r_i}: {cells}")
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
        print(f"  NOTES: {slide.notes_slide.notes_text_frame.text.strip()[:400]}")
