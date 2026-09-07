"""
Patch the Evidence & Artifacts slide with real, clickable GitHub links.

Edits the deck IN PLACE and locates the slide by its title text, so it survives
slides being added or deleted (e.g. the generic "How to Use" slide being removed).

The link boxes are ~2.1in wide, far too narrow for full GitHub URLs, so each shows
short display text with the real URL attached as a PowerPoint hyperlink.
"""

import sys
from pptx import Presentation
from pptx.util import Pt

DECK = sys.argv[1]
# Optional second argument: write elsewhere. Needed when the deck is open in
# PowerPoint, which holds a write lock on the file.
OUT = sys.argv[2] if len(sys.argv) > 2 else DECK
REPO = "https://github.com/shamirna-dev/dbt_capstone"
BRANCH = "capstone_L2"
TREE = f"{REPO}/tree/{BRANCH}"
BLOB = f"{REPO}/blob/{BRANCH}"

prs = Presentation(DECK)


def find_slide(title_text):
    for idx, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            if sh.has_text_frame and title_text.lower() in sh.text_frame.text.lower():
                return idx, slide
    raise SystemExit(f"could not find a slide titled '{title_text}'")


def set_link(shape, text, url, size=9):
    """Short display text, real URL behind it."""
    tf = shape.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    for r in list(para.runs[1:]):
        r._r.getparent().remove(r._r)
    run = para.runs[0] if para.runs else para.add_run()
    run.text = text
    run.font.size = Pt(size)
    if url:
        run.hyperlink.address = url


def set_label(shape, text, size=None):
    tf = shape.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    for r in list(para.runs[1:]):
        r._r.getparent().remove(r._r)
    run = para.runs[0] if para.runs else para.add_run()
    run.text = text
    if size:
        run.font.size = Pt(size)


idx, slide = find_slide("Evidence & Artifacts")
s = {sh.name: sh for sh in slide.shapes}
print(f"patching slide {idx}: Evidence & Artifacts")

# Labels retitled to describe what actually exists. There is no PR - the work is
# a branch - and the "demo recording" slot is used for the write-up.
set_label(s["Text 2"],  "\u2022  Repo & branch \u2014 all code")
set_label(s["Text 5"],  "\u2022  Write-up / design doc")
set_label(s["Text 8"],  "\u2022  Prompt log / agent transcript")
set_label(s["Text 11"], "\u2022  Dashboard source (screenshots above)")

set_link(s["Text 4"],  "dbt_capstone @ capstone_L2", TREE)
set_link(s["Text 7"],  "docs/CAPSTONE.md",           f"{BLOB}/docs/CAPSTONE.md")
set_link(s["Text 10"], "METRICS.AGENT_RUN_LOG \u00b7 sql/05",
         f"{BLOB}/sql/05_agent_framework.sql", size=8)
set_link(s["Text 13"], "app/streamlit_app.py",       f"{BLOB}/app/streamlit_app.py")

try:
    prs.save(OUT)
except PermissionError:
    raise SystemExit(
        f"cannot write {OUT} - it is open in PowerPoint. Close the deck and "
        f"re-run, or pass a second argument as the output path."
    )
print(f"SAVED: {OUT}")
